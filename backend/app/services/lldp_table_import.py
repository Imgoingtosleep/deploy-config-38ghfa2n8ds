"""
Import an LLDP neighbor table from Excel (.xlsx / .xlsm) or CSV and build the
topology from it.

Accepted tables:
  - The Excel exported here (LLDP_Inventory + Execution_Summary sheets)
  - Hand-made tables: column names are matched loosely (Local Device / Hostname /
    Switch, Local Port / Interface, Neighbor / System Name / Device ID, Port ID ...),
    the header may sit below a title, and blank Local Device cells (merged cells or
    one group per switch) repeat the value above.
  - One sheet per switch without a Local Device column: the sheet name is the device.
A sheet with devices but no neighbor column (e.g. Execution_Summary) adds the
devices' IP / model and devices that have no LLDP neighbor.
"""
import csv
import io
import re
from typing import Any, Dict, Iterable, List, Optional, Tuple

from app.services.lldp_service import LldpService

# Normalized header (lower case, letters / digits only) -> field
_ALIASES = {
    "local_device": [
        "localdevice", "localhostname", "localhost", "localname", "localsysname", "localsystemname", "localswitch",
        "device", "devicename", "hostname", "host", "sysname", "switch", "switchname", "sourcedevice", "source",
        "fromdevice", "node",
    ],
    "local_port": [
        "localport", "localinterface", "localintf", "localif", "interface", "intf", "port", "sourceport",
        "sourceinterface", "fromport", "frominterface",
    ],
    "local_ip": ["localip", "localipaddress", "localmgmtip", "ip", "ipaddress", "mgmtip", "managementip", "sourceip", "hostip"],
    "local_model": ["localmodel", "model", "sourcemodel", "devicemodel"],
    "remote_device": [
        "remotedevice", "remotehostname", "remotehost", "remotename", "remotesysname", "remotesystemname",
        "neighbor", "neighbour", "neighbordevice", "neighbourdevice", "neighbordev", "neighborname", "neighbourname",
        "neighborhostname", "systemname", "deviceid", "targetdevice", "target", "todevice", "peer", "peerdevice",
    ],
    "remote_port": [
        "remoteport", "remoteinterface", "remoteintf", "neighborport", "neighbourport", "neighborinterface",
        "neighbourinterface", "neighborintf", "portid", "targetport", "targetinterface", "toport", "tointerface",
        "peerport", "peerinterface",
    ],
    "remote_ip": [
        "remoteip", "remoteipaddress", "remotemgmtip", "neighborip", "neighbourip", "neighboripaddress",
        "managementaddress", "mgmtaddress", "targetip", "peerip",
    ],
    "remote_model": ["remotemodel", "neighbormodel", "neighbourmodel", "targetmodel", "peermodel"],
    "remote_desc": ["remotedescription", "systemdescription", "neighbordescription", "neighbourdescription", "sysdescr"],
    "local_type": ["localtype", "localrole", "localdevicetype", "type", "role", "devicetype", "sourcetype"],
    "remote_type": ["remotetype", "remoterole", "remotedevicetype", "neighbortype", "neighbourtype", "neighborrole", "targettype", "peertype"],
    "status": ["status"],
}
# Words people write in a Type column -> topology device type
_TYPE_WORDS = [
    ("firewall", ("firewall", "fw", "utm", "ngfw", "fortigate", "palo", "asa", "usg")),
    ("wireless", ("wireless", "ap", "accesspoint", "wifi", "wlan", "airengine")),
    ("router", ("router", "rt", "gateway", "gw", "wan")),
    ("server", ("server", "srv", "vm", "esxi")),
    ("cloud", ("cloud", "isp", "internet")),
    ("pc", ("pc", "host", "client", "workstation", "notebook", "laptop")),
    ("switch", ("switch", "sw", "l2", "l3", "core", "distribution", "access", "stack")),
]
_FIELD_BY_ALIAS = {alias: field for field, aliases in _ALIASES.items() for alias in aliases}
# Devices that did not answer in the exported Execution_Summary are not part of the network picture
_FAILED_STATUS_RE = re.compile(r"fail|unreachable|timeout|refused|duplicate", re.I)
_GENERIC_SHEET_RE = re.compile(r"^(sheet|แผ่นงาน|worksheet|lldp|data|table)\s*\d*$|inventory|summary|neighbou?rs?", re.I)
HEADER_SCAN_ROWS = 30
MAX_ROWS = 200000


class TableImportError(ValueError):
    pass


def _norm(header: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(header or "").lower())


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def device_type(value: str) -> str:
    """'Firewall', 'FW', 'L3 Switch', 'AP' ... -> router / switch / firewall / ...; '' when not understood"""
    v = (value or "").strip().lower()
    if not v:
        return ""
    words = set(re.split(r"[^a-z0-9]+", v)) | {re.sub(r"[^a-z0-9]", "", v)}
    for role, keys in _TYPE_WORDS:
        if any(k in words for k in keys):
            return role
    return ""


def _find_header(rows: List[List[str]]) -> Optional[Tuple[int, Dict[str, int]]]:
    """First row (within the first rows) whose cells name a device column; {field: column index}"""
    best = None
    for idx, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        cols: Dict[str, int] = {}
        for c, value in enumerate(row):
            field = _FIELD_BY_ALIAS.get(_norm(value))
            if field and field not in cols:
                cols[field] = c
        is_links = "remote_device" in cols and ("local_device" in cols or "local_port" in cols)
        is_hosts = "local_device" in cols and ("local_ip" in cols or "local_model" in cols)
        if is_links:
            return idx, cols
        if is_hosts and best is None:
            best = (idx, cols)
    return best


def _parse_sheet(name: str, rows: List[List[str]], neighbors: List[Dict[str, str]],
                 hosts: Dict[str, Dict[str, Any]], roles: Dict[str, str], warnings: List[str]) -> None:
    found = _find_header(rows)
    if not found:
        # Only tables are worth a warning; a notes sheet (one column of text) is not
        if any(sum(1 for c in r if c) >= 2 for r in rows[:HEADER_SCAN_ROWS]):
            warnings.append(f"Sheet '{name}': no Local Device / Neighbor columns found, skipped")
        return
    header_idx, cols = found
    get = lambda row, field: row[cols[field]] if field in cols and cols[field] < len(row) else ""  # noqa: E731

    if "remote_device" not in cols:
        # Device list only: IP / model and devices without neighbors
        for row in rows[header_idx + 1:]:
            name_ = get(row, "local_device")
            if not name_ or _FAILED_STATUS_RE.search(get(row, "status")):
                continue
            h = hosts.setdefault(name_, {"hostname": name_, "ip": "", "model": "", "rows": 0})
            h["ip"] = h["ip"] or get(row, "local_ip")
            h["model"] = h["model"] or get(row, "local_model")
            if device_type(get(row, "local_type")):
                roles.setdefault(name_, device_type(get(row, "local_type")))
        return

    sheet_device = ""
    if "local_device" not in cols:
        if _GENERIC_SHEET_RE.search(name.strip()):
            warnings.append(f"Sheet '{name}': no Local Device column and the sheet name is not a device name, skipped")
            return
        sheet_device = name.strip()

    last = {"local_device": sheet_device, "local_ip": "", "local_model": ""}
    skipped = 0
    for row in rows[header_idx + 1:]:
        if not any(row):
            continue
        local = get(row, "local_device") or sheet_device
        if local and local != last["local_device"]:
            last = {"local_device": local, "local_ip": get(row, "local_ip"), "local_model": get(row, "local_model")}
        elif not local:
            # Merged / grouped cells: the device written on an earlier row
            local = last["local_device"]
        remote = get(row, "remote_device")
        remote_port = get(row, "remote_port")
        if not local or not (remote or remote_port):
            skipped += 1
            continue
        local_ip = get(row, "local_ip") or last["local_ip"]
        local_model = get(row, "local_model") or last["local_model"]
        desc = get(row, "remote_desc")
        remote_model = get(row, "remote_model") or (LldpService.extract_model(desc) if desc else "")
        neighbors.append({
            "Local Device": local,
            "Local Model": local_model,
            "Local IP": local_ip,
            "Local Port": get(row, "local_port"),
            "Remote Device": remote or "N/A",
            "Remote Model": remote_model,
            "Remote Port": remote_port or "N/A",
            "Remote IP": get(row, "remote_ip"),
            "Remote Description": desc,
        })
        h = hosts.setdefault(local, {"hostname": local, "ip": "", "model": "", "rows": 0})
        h["ip"] = h["ip"] or local_ip
        h["model"] = h["model"] or local_model
        h["rows"] += 1
        for dev, field in ((local, "local_type"), (remote, "remote_type")):
            if dev and device_type(get(row, field)):
                roles.setdefault(dev, device_type(get(row, field)))
        if len(neighbors) > MAX_ROWS:
            raise TableImportError(f"More than {MAX_ROWS} LLDP rows")
    if skipped:
        warnings.append(f"Sheet '{name}': {skipped} row(s) without a device / neighbor were skipped")


def _sheets_from_xlsx(content: bytes) -> Iterable[Tuple[str, List[List[str]]]]:
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        raise TableImportError(f"Cannot read the Excel file ({type(e).__name__}: {e})")
    try:
        # Our own export: the inventory first, so its devices are known before the summary
        order = sorted(wb.worksheets, key=lambda ws: 0 if ws.title == "LLDP_Inventory" else 1)
        for ws in order:
            if ws.title == "Raw_Logs" or getattr(ws, "sheet_state", "visible") != "visible":
                continue
            rows = []
            for i, r in enumerate(ws.iter_rows(values_only=True)):
                if i > MAX_ROWS:
                    break
                rows.append([_cell(v) for v in r])
            yield ws.title, rows
    finally:
        wb.close()


def _rows_from_csv(text: str) -> List[List[str]]:
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [[_cell(v) for v in r] for r in csv.reader(io.StringIO(text), dialect)]


def looks_like_csv_table(text: str) -> bool:
    head = text[:4096].lower()
    return bool(re.search(r"(local|device|hostname|switch).{0,200}(remote|neighbou?r|system ?name|device ?id|port ?id)", head, re.S))


def parse_lldp_table(file_name: str, content: bytes, kind: str) -> Dict[str, Any]:
    """kind: 'xlsx' or 'csv'. Returns the same report shape as /lldp/discover (imported=True)"""
    neighbors: List[Dict[str, str]] = []
    hosts: Dict[str, Dict[str, Any]] = {}
    warnings: List[str] = []
    roles: Dict[str, str] = {}
    if kind == "xlsx":
        sheets = _sheets_from_xlsx(content)
    else:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("cp874", errors="replace")  # Thai Excel "CSV" without UTF-8
        sheets = [(re.sub(r"\.[^.]*$", "", file_name) or "CSV", _rows_from_csv(text))]
    for name, rows in sheets:
        _parse_sheet(name, rows, neighbors, hosts, roles, warnings)

    if not neighbors and not hosts:
        raise TableImportError(
            "No LLDP table found. The sheet needs a header row with at least a device column "
            "(Local Device / Hostname) and a neighbor column (Remote Device / Neighbor / System Name)."
        )

    host_rows = [
        {"hostname": h["hostname"], "ip": h["ip"], "model": h["model"], "success": True, "status": "IMPORTED"}
        for h in hosts.values()
    ]
    LldpService.fill_remote_models(neighbors, host_rows)
    topology = LldpService.build_topology(neighbors, host_rows)
    # A Type column in the table wins over the guess from the model
    if roles:
        for node in topology["nodes"]:
            node["role"] = roles.get(node["id"], node["role"])
        order = {"router": 0, "firewall": 1, "switch": 2, "wireless": 3, "server": 4, "unknown": 9}
        topology["nodes"].sort(key=lambda n: (order.get(n["role"], 5), -n["degree"], n["hostname"]))
    unknown = [n["hostname"] for n in topology["nodes"] if n["role"] == "unknown"]
    if unknown:
        warnings.append(
            f"{len(unknown)} device(s) have no model or type ({', '.join(unknown[:5])}{' …' if len(unknown) > 5 else ''}): "
            "add a Local Type / Remote Type column (Switch, Router, Firewall, Wireless AP ...) to set their icons"
        )
    counts = {h["hostname"]: h["rows"] for h in hosts.values()}
    report_hosts = [
        {
            "hostname": n["hostname"],
            "ip": n.get("ip", ""),
            "model": n.get("model", ""),
            "depth": 0,
            "status": "IMPORTED",
            "success": True,
            "error": None,
            "detail": f"{counts[n['id']]} LLDP row(s) in the table" if counts.get(n["id"]) else
            ("Listed in the table" if n["id"] in counts else "Seen as a neighbor only"),
            "neighbors_found": counts.get(n["id"], 0),
            "neighbors": [],
            "log": "",
            "raw_output": "",
            "execution_time_seconds": 0,
        }
        for n in topology["nodes"]
    ]
    return {
        "imported": True,
        "file_name": file_name,
        "format": kind,
        "total_hosts": len(report_hosts),
        "success_hosts": len(hosts),
        "failed_hosts": 0,
        "total_lldp_rows": len(neighbors),
        "overall_time_seconds": 0,
        "neighbors": neighbors,
        "hosts": report_hosts,
        "topology": topology,
        "annotations": {"flows": [], "zones": [], "notes": []},
        "import_warnings": warnings,
    }


def build_template() -> bytes:
    """Empty LLDP table with the recognized columns and two example rows"""
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LLDP_Inventory"
    headers = ["Local Device", "Local Model", "Local IP", "Local Port", "Remote Device", "Remote Model", "Remote Port", "Remote IP",
               "Local Type", "Remote Type"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    ws.append(["CORE-SW1", "S6730-H48X6C", "10.0.0.1", "XGE0/0/1", "ACC-SW1", "S5735-L48T4X-A1", "XGE0/0/49", "10.0.0.11", "Switch", "Switch"])
    ws.append(["", "", "", "XGE0/0/2", "ACC-SW2", "", "GE0/0/49", "10.0.0.12", "", "Switch"])
    ws.append(["CORE-SW1", "", "", "GE0/0/1", "FW-01", "FortiGate-100F", "port1", "10.0.0.254", "", "Firewall"])
    for i, width in enumerate([16, 18, 14, 14, 16, 18, 14, 14, 12, 12], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    notes = wb.create_sheet("How_to_use")
    for line in [
        "One row per LLDP neighbor. Only Local Device and Remote Device are required; ports make the links precise.",
        "A blank Local Device repeats the device of the row above (merged cells are fine).",
        "Other column names are understood too: Hostname / Switch, Interface, Neighbor / System Name / Device ID, Port ID ...",
        "Local Type / Remote Type (optional): Router, Switch, Firewall, Wireless AP, Server, Cloud, PC. Without them the model decides the icon.",
        "Models the page does not know: add Model Rules on the LLDP page, or fill in the Type columns.",
        "Import it with 'Import Topology' on the LLDP Discovery page.",
    ]:
        notes.append([line])
    notes.column_dimensions["A"].width = 120
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()
