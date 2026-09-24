import io
import ipaddress
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from typing import List, Dict, Any, Optional, Set

from app.core.config import settings
from app.schemas.device import DeviceCredentials
from app.services.netmiko_service import NetmikoService

SYSNAME_CMD = "display current-configuration | include sysname"
LLDP_BRIEF_CMD = "display lldp neighbor brief"
LLDP_DETAIL_CMD = "display lldp neighbor interface {intf}"
LLDP_FULL_CMD = "display lldp neighbor"
VERSION_CMD = "display version"

# Huawei product models, e.g. S2750-28TP-EI-AC, CE6881-48S6CQ, AR6120-S, NE40E-X8A.
# Longest match wins so "S5720-28X-SI-AC" beats the short "(S5720 V200R011...)".
_MODEL_RE = re.compile(r"(?<![\w-])((?:AirEngine|ATN|USG|CE|CX|AR|NE|AC|S)\d{2,5}[A-Z]*(?:-[A-Z0-9]+)*)(?!\w)")
ROUTER_MODEL_PREFIXES = ("AR", "NE")

# Excel cell hard limit is 32,767 chars
EXCEL_CELL_LIMIT = 32000

_BRIEF_HEADER_ALIASES = {
    "local": ["Local Interface", "Local Intf", "Local Port"],
    "remote_dev": ["Neighbor Device", "Neighbor Dev", "System Name"],
    "remote_intf": ["Neighbor Interface", "Neighbor Intf", "Neighbor Port", "Port ID"],
    "exptime": ["Exptime(s)", "Exptime", "Expire"],
}

# Cisco 'show lldp neighbors' puts Device ID first, then Local Intf ... Port ID
_BRIEF_HEADER_ALIASES_CISCO = {
    "remote_dev": ["Device ID"],
    "local": ["Local Intf", "Local Interface"],
    "remote_intf": ["Port ID"],
    "exptime": ["Hold-time"],
}

# Raisecom ROS 'show lldp remote':
#   Port      ChassisId        PortId       SysName     MgtAddress     ExpiredTime
_BRIEF_HEADER_ALIASES_RAISECOM = {
    "local": ["Local Port", "Local Intf", "Interface", "Port"],
    "chassis": ["ChassisId", "Chassis ID", "Chassis Id"],
    "remote_intf": ["PortId", "Port ID", "Port Id"],
    "remote_dev": ["SysName", "System Name"],
    "mgmt": ["MgtAddress", "Management Address", "MgmtAddress", "Mgmt Address"],
    "exptime": ["ExpiredTime", "Expired Time", "TTL"],
}

# The device parser rejected the command, so this command profile is the wrong vendor
# ('% Unknown command.' is Raisecom ROS)
_CLI_REJECT_RE = re.compile(
    r"unrecognized command|invalid input|% invalid|syntax error|wrong parameter|incomplete command|unknown command",
    re.I,
)

# Cisco model codes, e.g. WS-C2960X-48TS-L, C9200-48P, ISR4331, N9K-C93180YC-EX, IE-3300-8T2S
_MODEL_RE_CISCO = re.compile(
    r"(?<![\w-])((?:WS-C|N\dK-C|IE-|ISR|ASR|CSR|CISCO|CBS|SG|C)\d{3,5}[A-Z0-9]*(?:-[A-Z0-9]+)*)(?![\w(])"
)
# 'Model number : WS-C2960X-48TS-L' (show version) / 'cisco C9300-48P (X86) processor'
_MODEL_LINE_CISCO = re.compile(r"^\s*Model\s*number\s*:\s*(\S+)", re.MULTILINE | re.IGNORECASE)
_MODEL_HW_CISCO = re.compile(r"^\s*cisco\s+(?:Nexus\S*\s+)?(\S+)\s.*?(?:processor|chassis)", re.MULTILINE | re.IGNORECASE)
# IOS image names look like models but are not: C2960X-UNIVERSALK9-M, C3750E-IPBASEK9-M
_IMAGE_TOKEN_RE = re.compile(r"K9|UNIVERSAL|IPBASE|LANBASE|IPSERVICES|ENTSERVICES|ADVIP|^M$|^MZ$", re.I)
# 'Catalyst L3 Switch Software (CAT9K_IOSXE)' carries only the family
_FAMILY_RE_CISCO = re.compile(r"\bCAT(\d{1,2})K(?![A-Za-z0-9])", re.I)
_CISCO_ROUTER_RE = re.compile(r"^(?:ISR|ASR|CSR|CISCO\d{4}|C8\d{3}|C11\d{2}|C[1-3]900$)", re.I)

# Raisecom product models, e.g. ISCOM2600G-4GE-AC, ISCOM2128EA-MA, RAX711-C-AC, iTN201, RC002-16
_MODEL_RE_RAISECOM = re.compile(
    r"(?<![\w-])((?:ISCOM|RAX|iTN|RC|Gazelle)\d{2,5}[A-Z0-9]*(?:-[A-Z0-9]+)*)(?![\w(])",
    re.I,
)

# Juniper model regexes
_MODEL_LINE_JUNIPER = re.compile(r"^\s*Model\s*:\s*([A-Za-z0-9_+-]+)", re.MULTILINE | re.IGNORECASE)
_MODEL_RE_JUNIPER = re.compile(r"(?<![\w-])((?:SRX|QFX|EX|MX|ACX|PTX|NFX)\d{3,5}[A-Z0-9]*(?:-[A-Z0-9+]+)*)(?!\w)", re.I)

# Aruba / HP ProCurve model regexes
_MODEL_LINE_ARUBA = re.compile(r"MODEL:\s*([A-Za-z0-9_+-]+)", re.I)
_MODEL_RE_ARUBA = re.compile(
    r"(?<![\w-])((?:(?:CX\s*)?[1-9]\d{3}[A-Z]*(?:-[A-Z0-9+]+)*|JL\d{3}[A-Z]|J\d{4}[A-Z]|Aruba\d{3,4}[A-Z0-9-]*))(?!\w)",
    re.I,
)

# HP / H3C Comware model regexes
_MODEL_RE_HP = re.compile(
    r"(?<![\w-])((?:S\d{4}[A-Z0-9]*(?:-[A-Z0-9+]+)*|MSR\d{4}[A-Z0-9-]*|\d{4}(?:AF|X)?(?:-[A-Z0-9+]+)+|HPE\s+\d{4}[A-Z0-9\s+-]*EI|HPE\s+\d{4}))(?!\w)",
    re.I,
)

# MikroTik RouterOS model regexes
_MODEL_LINE_MIKROTIK = re.compile(r"^\s*(?:board-name|model)\s*:\s*(.+)$", re.MULTILINE | re.IGNORECASE)
_MODEL_RE_MIKROTIK = re.compile(
    r"(?<![\w-])((?:CCR|CRS|CSS|RB|hEX|hAP|NetMetal|PowerBox)\d{2,5}[A-Z0-9]*(?:-[A-Z0-9+]+)*)(?!\w)",
    re.I,
)


# Netmiko driver to log in with for each command profile parser: an LLDP neighbor
# has no device type of its own, so its driver is swept in command profile order
PARSER_DRIVERS = {
    "huawei": "huawei",
    "cisco": "cisco_ios",
    "raisecom": "raisecom_roap",
    "aruba": "aruba_os",
    "hp_comware": "hp_comware",
    "juniper": "juniper_junos",
    "mikrotik": "mikrotik_routeros",
}


def cli_rejected(text: str) -> bool:
    """True when the CLI answered with a syntax error instead of running the command"""
    return bool(text) and bool(_CLI_REJECT_RE.search(text))


class LldpService:
    @classmethod
    def name_key(cls, name: str) -> str:
        """Device name as it is compared: domain dropped, case-insensitive"""
        return cls.short_hostname(name).strip().lower()

    @staticmethod
    def short_hostname(name: str) -> str:
        """
        Cisco with 'ip domain-name lab.local' advertises the LLDP System Name /
        Device ID as 'router_cisco.lab.local' (the brief table even cuts it to
        'router_cisco.lab.loc'). Drop the domain so it matches the SSH'd hostname.
        IP addresses and names like 'SW-1.2F' are left alone.
        """
        # rstrip: the brief table can cut the name right after the dot ('router_cisco.')
        name = (name or "").strip().rstrip(".")
        if not name or re.fullmatch(r"\d{1,3}(?:\.\d{1,3}){3}", name):
            return name
        host, dot, domain = name.partition(".")
        # Any 'ip domain-name' value (lab.local, corp.example.com, lab1, net-2, or a
        # truncated '.c'): DNS labels whose first label starts with a letter.
        # 'SW-1.2F' / 'R1.1' (digit right after the dot) are kept as real names.
        if dot and host and re.fullmatch(r"[A-Za-z][A-Za-z0-9-]*(?:\.[A-Za-z0-9-]*)*", domain):
            return host
        return name

    @classmethod
    def get_sysname(cls, prompt: str) -> str:
        return cls.short_hostname(re.sub(r"[<>\[\]#]", "", prompt or "").strip())

    @staticmethod
    def detect_vendor(text: str, parser: str = "") -> str:
        """Vendor of the device the text describes; falls back to the command profile's parser"""
        t = (text or "").lower()
        if "juniper" in t or "junos" in t:
            return "juniper"
        if "aruba" in t or "procurve" in t:
            return "aruba"
        if "h3c" in t or "comware" in t or "hpe" in t:
            return "hp_comware"
        if "mikrotik" in t or "routeros" in t or "routerboard" in t or "crs" in t or "ccr" in t:
            return "mikrotik"
        if "raisecom" in t:
            return "raisecom"
        if "huawei" in t or "vrp" in t:
            return "huawei"
        if "cisco" in t or "nx-os" in t or "ios-xe" in t:
            return "cisco"
        # Raisecom descriptions without the company name: "ROS_5.2.1", "ISCOM2608G", "RAX711".
        # Whole words only - the substrings hit "microsoft", "syntax", "across"
        if re.search(r"(?<![a-z])ros(?![a-z])|\biscom\d|\brax\d", t):
            return "raisecom"
        return parser or "huawei"

    @staticmethod
    def _extract_cisco_model(text: str) -> str:
        for rx in (_MODEL_LINE_CISCO, _MODEL_HW_CISCO):
            m = rx.search(text)
            if m and _MODEL_RE_CISCO.fullmatch(m.group(1)):
                return m.group(1)
        full, image = [], []
        for cand in _MODEL_RE_CISCO.findall(text):
            parts = cand.split("-")
            keep = []
            for part in parts:
                if _IMAGE_TOKEN_RE.search(part):
                    break
                keep.append(part)
            if len(keep) == len(parts):
                full.append(cand)
            elif keep:
                # C2960X-UNIVERSALK9-M -> C2960X (the platform family)
                image.append("-".join(keep))
        if full:
            return max(full, key=len)
        if image:
            return max(image, key=len)
        fam = _FAMILY_RE_CISCO.search(text)
        return f"C{fam.group(1)}K" if fam else ""

    @staticmethod
    def _extract_huawei_model(text: str) -> str:
        candidates = _MODEL_RE.findall(text)
        return max(candidates, key=len) if candidates else ""

    @staticmethod
    def _extract_juniper_model(text: str) -> str:
        m = _MODEL_LINE_JUNIPER.search(text)
        if m:
            return m.group(1).strip()
        cands = _MODEL_RE_JUNIPER.findall(text)
        return max(cands, key=len) if cands else ""

    @staticmethod
    def _extract_aruba_model(text: str) -> str:
        m = _MODEL_LINE_ARUBA.search(text)
        if m:
            return m.group(1).strip()
        cands = _MODEL_RE_ARUBA.findall(text)
        if not cands:
            return ""
        # Prefer switch series name (e.g. 2930F-24G-4SFP+, 6300M, 2530-24G-PoEP) over bare part number (JL665A, J9773A)
        series_cands = [c for c in cands if not re.fullmatch(r"J[L\d]\d{3}[A-Z]", c, re.I)]
        if series_cands:
            return max(series_cands, key=len)
        return max(cands, key=len)

    @staticmethod
    def _extract_hp_model(text: str) -> str:
        cands = _MODEL_RE_HP.findall(text)
        return max(cands, key=len).strip() if cands else ""

    @staticmethod
    def _extract_mikrotik_model(text: str) -> str:
        m = _MODEL_LINE_MIKROTIK.search(text)
        if m:
            return m.group(1).strip()
        cands = _MODEL_RE_MIKROTIK.findall(text)
        return max(cands, key=len) if cands else ""

    @staticmethod
    def _extract_raisecom_model(text: str) -> str:
        cands = _MODEL_RE_RAISECOM.findall(text)
        return max(cands, key=len) if cands else ""

    # ---- Command profile regexes: run on one command's output before the parser sees it ----
    @staticmethod
    def brief_header(lines: List[str], parser: str = "huawei"):
        """Index and column positions of the neighbor table header, (None, []) when absent"""
        aliases_map = {
            "cisco": _BRIEF_HEADER_ALIASES_CISCO,
            "raisecom": _BRIEF_HEADER_ALIASES_RAISECOM,
        }.get(parser, _BRIEF_HEADER_ALIASES)
        for idx, line in enumerate(lines):
            found = []
            for key, aliases in aliases_map.items():
                for alias in aliases:
                    pos = line.find(alias)
                    if pos >= 0:
                        found.append((pos, key))
                        break
            if any(k == "local" for _, k in found) and len(found) >= 2:
                return idx, sorted(found)
        return None, []

    @classmethod
    def keep_brief_header(cls, raw: str, kept: str, parser: str) -> str:
        """The brief parser takes its column positions from the table header, so a regex that
        keeps only neighbor rows must not throw that header away"""
        if cls.brief_header(kept.splitlines(), parser)[0] is not None:
            return kept
        idx, _ = cls.brief_header(raw.splitlines(), parser)
        return f"{raw.splitlines()[idx]}\n{kept}" if idx is not None else kept

    @staticmethod
    def regex_capture(text: str, pattern: str) -> Optional[str]:
        """
        Value read by a command profile regex: capture group 1 when the pattern has one,
        otherwise the whole match. None when the pattern is empty, does not compile or does
        not match, and the caller then falls back to the built-in parsing.
        """
        if not pattern or not text:
            return None
        try:
            m = re.search(pattern, text, re.MULTILINE | re.IGNORECASE)
        except re.error:
            return None
        if not m:
            return None
        value = (m.group(1) if m.groups() else m.group(0)) or ""
        return value.strip() or None

    @staticmethod
    def regex_lines(text: str, pattern: str) -> Optional[str]:
        """Only the lines a command profile regex matches, None when the pattern is empty,
        does not compile or keeps nothing (the caller then parses the untouched output)"""
        if not pattern or not text:
            return None
        try:
            rx = re.compile(pattern, re.IGNORECASE)
        except re.error:
            return None
        kept = [line for line in text.splitlines() if rx.search(line)]
        return "\n".join(kept) if kept else None

    @classmethod
    def extract_model(cls, text: str, parser: str = "") -> str:
        """
        Model from 'System description' or version output, '' when none is found.
        The vendor is taken from the text itself, so a Huawei switch still reads a
        Cisco neighbor's model (and vice versa); the parser is only a fallback.
        """
        # Model rules edited on the LLDP page win over the built-in regexes
        from app.services.model_rule_service import ModelRuleService
        custom = ModelRuleService.match_model(text or "")
        return custom["model"] if custom else cls.builtin_model(text, parser)

    @classmethod
    def builtin_model(cls, text: str, parser: str = "") -> str:
        """Model from the built-in Huawei, Cisco, Raisecom, Juniper, Aruba, HP, MikroTik regexes"""
        text = text or ""
        vendor = cls.detect_vendor(text, parser)
        if vendor == "juniper":
            return cls._extract_juniper_model(text)
        if vendor == "aruba":
            return cls._extract_aruba_model(text)
        if vendor == "hp_comware":
            return cls._extract_hp_model(text)
        if vendor == "mikrotik":
            return cls._extract_mikrotik_model(text)
        if vendor == "raisecom":
            return cls._extract_raisecom_model(text)
        if vendor == "cisco":
            return cls._extract_cisco_model(text)
        if vendor == "huawei":
            hw = cls._extract_huawei_model(text)
            if hw:
                return hw

        # Cascading fallback across all vendor extractors
        for extractor in [
            cls._extract_huawei_model,
            cls._extract_cisco_model,
            cls._extract_juniper_model,
            cls._extract_aruba_model,
            cls._extract_raisecom_model,
            cls._extract_hp_model,
            cls._extract_mikrotik_model,
        ]:
            cand = extractor(text)
            if cand:
                return cand
        return ""

    @classmethod
    def device_role(cls, model: str) -> str:
        """A model rule with a role wins over the built-in router / switch guess"""
        return cls.device_role_info(model)["role"]

    @classmethod
    def device_role_info(cls, model: str) -> Dict[str, Any]:
        """
        The role and where it came from, so the UI can show which model is drawn
        with which icon and why: {role, source: 'rule'|'builtin'|'none', rule_id, rule_name}
        """
        from app.services.model_rule_service import ModelRuleService
        custom = ModelRuleService.match_role(model)
        if custom:
            rule = custom["rule"]
            return {
                "role": custom["role"],
                "source": "rule",
                "rule_id": rule.get("id", ""),
                "rule_name": rule.get("name", ""),
            }
        role = cls.builtin_role(model)
        return {
            "role": role,
            "source": "none" if role == "unknown" else "builtin",
            "rule_id": "",
            "rule_name": "",
        }

    @staticmethod
    def builtin_role(model: str) -> str:
        """Determine device role (router, switch, firewall, wireless) from model string"""
        if not model:
            return "unknown"
        m = model.upper().strip()
        # Firewalls
        if m.startswith(("SRX", "USG", "ASA", "FTD", "FORTIGATE")):
            return "firewall"
        # Wireless APs
        if m.startswith(("AIR-", "AP-", "AIRENGINE")):
            return "wireless"
        # Routers
        if (
            (m.startswith(ROUTER_MODEL_PREFIXES) and not m.startswith("ARUBA")) # AR, NE
            or _CISCO_ROUTER_RE.match(m)       # ISR, ASR, CSR, C8xx, C11xx, C1-3900
            or m.startswith(("RAX", "MX", "ACX", "PTX", "MSR", "CCR", "HEX", "HAP", "RB"))
        ):
            return "router"
        # Switches (all remaining models like Catalyst, S-series, ISCOM, QFX, EX, ProCurve, CRS, CSS, etc.)
        return "switch"

    @staticmethod
    def fill_remote_models(neighbors: List[Dict[str, Any]], hosts: List[Dict[str, Any]]) -> None:
        """Blank 'Remote Model' cells get the model the neighbor reported about itself when it was SSH'd"""
        by_name = {h["hostname"]: h["model"] for h in hosts if h.get("model")}
        by_ip = {h["ip"]: h["model"] for h in hosts if h.get("model") and h.get("ip")}
        for n in neighbors:
            if not n.get("Remote Model"):
                n["Remote Model"] = by_name.get(n.get("Remote Device"), "") or by_ip.get(n.get("Remote IP"), "")

    @staticmethod
    def parse_lldp_brief(output: str, parser: str = "huawei") -> List[Dict[str, str]]:
        """
        Parse an LLDP neighbor table into rows of {local_port, remote_device,
        remote_port}. Column boundaries are taken from the header line, so both
        Huawei 'display lldp neighbor brief' and Cisco 'show lldp neighbors'
        work, as do column order differences between firmware versions.
        """
        rows: List[Dict[str, str]] = []
        lines = (output or "").splitlines()
        header_idx, columns = LldpService.brief_header(lines, parser)
        if header_idx is None:
            return rows

        for line in lines[header_idx + 1:]:
            if not line.strip() or set(line.strip()) <= set("-="):
                continue
            if re.match(r"^\s*[<\[]\S+[>\]]\s*$", line):
                continue
            values = {}
            for i, (start, key) in enumerate(columns):
                end = columns[i + 1][0] if i + 1 < len(columns) else len(line)
                # First column starts at 0 so leading data is not lost
                values[key] = line[0 if i == 0 else start:end].strip()
            local = values.get("local", "")
            if not local or " " in local:
                # Fallback for misaligned rows: whitespace split in header order
                parts = line.split()
                if len(parts) < len(columns):
                    continue
                values = {key: parts[i] for i, (_, key) in enumerate(columns)}
                local = values.get("local", "")
            # Interface names like GE0/0/1, 10GE1/0/48, Eth-Trunk1, MEth0/0/1
            if not re.match(r"^(?=[\w\-]*[A-Za-z])[\w\-]+\d", local):
                continue
            mgmt = re.search(r"\d{1,3}(?:\.\d{1,3}){3}", values.get("mgmt", ""))
            rows.append({
                "local_port": local,
                "remote_device": LldpService.short_hostname(values.get("remote_dev", "")),
                "remote_port": values.get("remote_intf", ""),
                # Only tables with a management address column (Raisecom 'show lldp remote')
                "remote_ip": mgmt.group(0) if mgmt else "",
            })
        return rows

    @staticmethod
    def intf_key(name: str) -> str:
        """Match short and long interface names: GE0/0/1 == GigabitEthernet0/0/1, XGE0/0/1 == XGigabitEthernet0/0/1"""
        s = re.sub(r"\s+", "", name or "")
        m = re.match(r"^(\d*[A-Za-z])[A-Za-z\-]*?(\d+(?:/\d+)*(?:[:.]\d+)?)$", s)
        if not m:
            return s.lower()
        prefix = m.group(1).upper()
        if "trunk" in s.lower():
            prefix += "TRUNK"
        return f"{prefix}{m.group(2)}"

    @staticmethod
    def extract_mgmt_ipv4(text: str) -> str:
        """
        First valid IPv4 from 'Management address[ value] : <value>' lines.
        The value must be on the same line (an empty value must not pick up the
        next line), and MAC / IPv6 management addresses are skipped.
        """
        for m in re.finditer(r"^[ \t]*Management address(?:[ \t]+value)?[ \t]*:[ \t]*(\S*)", text or "", re.MULTILINE | re.IGNORECASE):
            candidate = re.match(r"\d{1,3}(?:\.\d{1,3}){3}(?![\d.])", m.group(1))
            if not candidate:
                continue
            try:
                return str(ipaddress.IPv4Address(candidate.group(0)))
            except ValueError:
                continue
        return ""

    @classmethod
    def parse_lldp_detail(cls, output: str, default_local_port: Optional[str] = None) -> List[Dict[str, str]]:
        """
        Parse Huawei 'display lldp neighbor [interface X]' detail output.
        Same block split as the original script, plus support for several
        neighbors per interface and the management address. When the
        '<port> has N neighbor(s):' header is missing (some VRP versions on
        the per-interface command), the whole output is used for default_local_port.
        """
        results: List[Dict[str, str]] = []
        blocks = re.split(r"(\S+) has \d+ neighbors?(?:\(s\))?\s*:", output or "", flags=re.IGNORECASE)
        if len(blocks) < 3 and default_local_port:
            blocks = ["", default_local_port, output or ""]
        for i in range(1, len(blocks), 2):
            local_port = blocks[i]
            content = blocks[i + 1] if i + 1 < len(blocks) else ""
            neighbors = re.split(r"Neighbor index\s*:\s*\d+", content, flags=re.IGNORECASE)
            neighbors = [n for n in neighbors if n.strip()] or [content]
            for n in neighbors:
                sysname = re.search(r"^\s*System\s*name\s*:(.*)$", n, re.MULTILINE | re.IGNORECASE)
                port_id = re.search(r"^\s*Port\s*ID\s*:(.*)$", n, re.MULTILINE | re.IGNORECASE)
                if not sysname and not port_id:
                    continue
                remote_ip = cls.extract_mgmt_ipv4(n)
                # System description spans several lines, up to the next 'Key :' line
                desc = re.search(
                    r"^\s*System\s*description\s*:(.*?)(?=^\s*[A-Za-z][A-Za-z /()\-]{2,40}:|\Z)",
                    n, re.MULTILINE | re.IGNORECASE | re.DOTALL,
                )
                results.append({
                    "local_port": local_port,
                    "remote_device": cls.short_hostname(sysname.group(1)) if sysname else "N/A",
                    "remote_port": port_id.group(1).strip() if port_id else "N/A",
                    "remote_ip": remote_ip,
                    "remote_model": cls.extract_model(desc.group(1), "huawei") if desc else "",
                    "remote_description": desc.group(1).strip() if desc else "",
                })
        return results

    @classmethod
    def parse_lldp_detail_cisco(cls, output: str, default_local_port: Optional[str] = None) -> List[Dict[str, str]]:
        """
        Parse Cisco 'show lldp neighbors [<intf>] detail'. Entries are separated
        by dashed rules; each block carries Local Intf / Port id / System Name /
        System Description and an indented 'IP:' under Management Addresses.
        """
        results: List[Dict[str, str]] = []
        blocks = [b for b in re.split(r"^-{4,}\s*$", output or "", flags=re.MULTILINE) if b.strip()]
        for block in blocks:
            sysname = re.search(r"^\s*System Name\s*:(.*)$", block, re.MULTILINE | re.IGNORECASE)
            port_id = re.search(r"^\s*Port id\s*:(.*)$", block, re.MULTILINE | re.IGNORECASE)
            if not sysname and not port_id:
                continue
            local = re.search(r"^\s*Local Intf\s*:(.*)$", block, re.MULTILINE | re.IGNORECASE)
            local_port = local.group(1).strip() if local else (default_local_port or "")
            if not local_port:
                continue
            ip_match = re.search(r"^\s*IP\s*:\s*(\d{1,3}(?:\.\d{1,3}){3})\s*$", block, re.MULTILINE | re.IGNORECASE)
            remote_ip = ""
            if ip_match:
                try:
                    remote_ip = str(ipaddress.IPv4Address(ip_match.group(1)))
                except ValueError:
                    remote_ip = ""
            desc = re.search(
                r"^\s*System Description\s*:(.*?)(?=^\s*(?:Time remaining|System Capabilities|Enabled Capabilities|Management Addresses|Auto Negotiation)\s*:|\Z)",
                block, re.MULTILINE | re.IGNORECASE | re.DOTALL,
            )
            results.append({
                "local_port": local_port,
                "remote_device": cls.short_hostname(sysname.group(1)) if sysname else "N/A",
                "remote_port": port_id.group(1).strip() if port_id else "N/A",
                "remote_ip": remote_ip,
                "remote_model": cls.extract_model(desc.group(1), "cisco") if desc else "",
                "remote_description": desc.group(1).strip() if desc else "",
            })
        return results

    @classmethod
    def parse_lldp_detail_raisecom(cls, output: str, default_local_port: Optional[str] = None) -> List[Dict[str, str]]:
        """
        Parse Raisecom ROS 'show lldp remote [interface <if>] detail'. One block per
        neighbor, headed by its local port ('Port gigaethernet1/1/1:', 'Port 1:',
        'Interface GE1/1/1 has 1 neighbor(s):'), then 'Key : value' lines with the ROS
        names (ChassisId, PortId, SysName, SysDesc, MgtAddress) or the long ones
        (Port ID, System Name, System Description, Management Address).
        """
        results: List[Dict[str, str]] = []
        # A block header is a line naming the local port, not a 'PortId :' / 'PortDesc :' field
        header = re.compile(
            r"^[ \t]*(?:Local[ \t]*)?(?:Port|Interface|Intf)(?![ \t]*(?:Id|Desc|Description|Subtype|Name)\b)"
            r"[ \t:]+(\S[^\r\n]*?)[ \t]*(?:has[ \t]+\d+[ \t]+neighbors?(?:\(s\))?)?[ \t]*:?[ \t]*$",
            re.MULTILINE | re.IGNORECASE,
        )
        parts = header.split(output or "")
        blocks = [(parts[i].strip(), parts[i + 1]) for i in range(1, len(parts), 2)]
        if not blocks and default_local_port:
            blocks = [(default_local_port, output or "")]
        for local_port, content in blocks:
            sysname = re.search(r"^\s*Sys(?:tem)?\s*Name\s*:(.*)$", content, re.MULTILINE | re.IGNORECASE)
            port_id = re.search(r"^\s*Port\s*Id\s*:(.*)$", content, re.MULTILINE | re.IGNORECASE)
            if not sysname and not port_id:
                continue
            remote_ip = ""
            for m in re.finditer(r"^\s*M(?:gt|gmt|anagement)\s*Addr(?:ess)?[^:\r\n]*:(.*)$", content, re.MULTILINE | re.IGNORECASE):
                ip = re.search(r"(?<![\d.])\d{1,3}(?:\.\d{1,3}){3}(?![\d.])", m.group(1))
                if ip:
                    try:
                        remote_ip = str(ipaddress.IPv4Address(ip.group(0)))
                        break
                    except ValueError:
                        continue
            # System description spans several lines, up to the next 'Key :' line
            desc = re.search(
                r"^\s*Sys(?:tem)?\s*Desc(?:ription)?\s*:(.*?)(?=^\s*[A-Za-z][A-Za-z /()\-]{2,40}:|\Z)",
                content, re.MULTILINE | re.IGNORECASE | re.DOTALL,
            )
            results.append({
                "local_port": local_port or default_local_port or "",
                "remote_device": cls.short_hostname(sysname.group(1)) if sysname else "N/A",
                "remote_port": port_id.group(1).strip() if port_id else "N/A",
                "remote_ip": remote_ip,
                "remote_model": cls.extract_model(desc.group(1), "raisecom") if desc else "",
                "remote_description": desc.group(1).strip() if desc else "",
            })
        return results

    @classmethod
    def parse_detail(cls, output: str, parser: str, default_local_port: Optional[str] = None) -> List[Dict[str, str]]:
        if parser == "cisco":
            return cls.parse_lldp_detail_cisco(output, default_local_port=default_local_port)
        if parser == "raisecom":
            return cls.parse_lldp_detail_raisecom(output, default_local_port=default_local_port)
        return cls.parse_lldp_detail(output, default_local_port=default_local_port)

    @classmethod
    def collect_device(
        cls,
        device: DeviceCredentials,
        depth: int = 0,
        command_profile_ids: Optional[List[str]] = None,
        driver: Optional[str] = None,
        cmd_profiles: Optional[List[Dict[str, Any]]] = None,
    ) -> Dict[str, Any]:
        """
        SSH to one device -> 'display lldp neighbor brief' -> loop every local
        interface that has a neighbor with 'display lldp neighbor interface <if>'.

        driver:       netmiko device type to log in with (chosen by collect_device_sweep)
        cmd_profiles: command profiles to run, already resolved (the sweep passes
                      only the ones whose parser matches the driver)
        """
        start = time.time()
        local_ip = device.host or ""
        log_lines = [f"--- Execution Log for {device.name or local_ip} ({local_ip}) ---"]
        raw_parts: List[str] = []
        neighbors: List[Dict[str, Any]] = []
        sysname = device.name or local_ip
        model = ""
        status = "Success"
        error = None
        used_profile = ""
        version_output = ""

        from app.services.command_profile_service import CommandProfileService
        if cmd_profiles is None:
            cmd_profiles = CommandProfileService.resolve_ordered(command_profile_ids)
        if driver:
            device.device_type = driver
            log_lines.append(f"Step 0a: SSH driver [{driver}]")
        rejected_profiles = 0

        try:
            with NetmikoService.connect_with_fallback(device) as (net_connect, winning_cred, attempt_logs):
                log_lines.append(f"Step 0: Connected via {winning_cred}")

                # Walk the command profiles on this one session: Huawei commands
                # first, and when the CLI rejects them retry with the next profile
                # (Cisco) instead of logging in again.
                for prof_no, cprof in enumerate(cmd_profiles, 1):
                    is_last_profile = prof_no == len(cmd_profiles)
                    cmds = cprof["commands"]
                    parser = cprof["parser"]
                    rxs = cprof.get("regexes") or {}
                    used_profile = cprof["name"]

                    def keep_lines(raw: str, field: str, step: str) -> str:
                        """Trim one command's output with this profile's regex, if it has one"""
                        pattern = rxs.get(field) or ""
                        if not pattern:
                            return raw
                        kept = cls.regex_lines(raw, pattern)
                        if kept is not None and field == "lldp_brief":
                            kept = cls.keep_brief_header(raw, kept, parser)
                        if kept is None:
                            log_lines.append(f"{step}: regex for '{field}' matched no line, parsing the full output")
                            return raw
                        log_lines.append(
                            f"{step}: regex for '{field}' kept {len(kept.splitlines())} of "
                            f"{len(raw.splitlines())} line(s)"
                        )
                        return kept

                    log_lines.append(
                        f"Step 0b: Command profile {prof_no}/{len(cmd_profiles)} [{cprof['name']}] (parser: {parser})"
                    )
                    neighbors = []
                    raw_parts.append(f"===== Command profile: {cprof['name']} =====")

                    if cmds.get("pager_disable"):
                        try:
                            net_connect.send_command(cmds["pager_disable"], read_timeout=settings.DEFAULT_TIMEOUT)
                        except Exception:
                            pass
                    sysname_source = cmds.get("sysname") or "prompt"
                    cfg_sysname = ""
                    try:
                        sys_raw = NetmikoService.clean_cli_output(
                            net_connect.send_command(cmds["sysname"], read_timeout=settings.DEFAULT_TIMEOUT)
                        )
                        raw_parts.append(f"<{sysname}> {cmds['sysname']}\n{sys_raw}")
                        custom_sysname = cls.regex_capture(sys_raw, rxs.get("sysname") or "")
                        if custom_sysname:
                            cfg_sysname = cls.short_hostname(custom_sysname.strip('"'))
                            sysname_source = f"regex on '{cmds['sysname']}'"
                        else:
                            if rxs.get("sysname"):
                                log_lines.append("Step 1: regex for 'sysname' did not match, using the built-in pattern")
                            # [ \t] instead of \s: a bare 'hostname' line must not swallow the next line
                            # of config (e.g. 'ip domain-name lab.local') as the device name
                            m = re.search(r"^[ \t]*(?:sysname|hostname)[ \t]+(\S[^\r\n]*?)[ \t]*$", sys_raw, re.MULTILINE | re.IGNORECASE)
                            cfg_sysname = cls.short_hostname(m.group(1).strip().strip('"')) if m else ""
                    except Exception as e:
                        log_lines.append(f"Step 1: '{cmds.get('sysname')}' failed: {e}")
                    if cfg_sysname:
                        sysname = cfg_sysname
                    else:
                        # Fall back to the CLI prompt when sysname is not in the config output
                        sysname_source = "prompt"
                        sysname = cls.get_sysname(net_connect.find_prompt()) or sysname
                    log_lines.append(f"Step 1: Identified real sysname as [{sysname}] (from {sysname_source})")

                    ver_raw = ""
                    try:
                        ver_raw = NetmikoService.clean_cli_output(
                            net_connect.send_command(cmds["version"], read_timeout=settings.DEFAULT_TIMEOUT)
                        )
                        raw_parts.append(f"<{sysname}> {cmds['version']}\n{ver_raw}")
                        custom_model = cls.regex_capture(ver_raw, rxs.get("version") or "")
                        if custom_model:
                            model = custom_model
                            model_source = f"regex on '{cmds['version']}'"
                        else:
                            if rxs.get("version"):
                                log_lines.append("Step 1b: regex for 'version' did not match, using the model rules")
                            model = cls.extract_model(ver_raw, parser)
                            model_source = f"'{cmds['version']}'"
                        version_output = ver_raw
                        log_lines.append(f"Step 1b: Model [{model or 'unknown'}] ({cls.device_role(model)}) from {model_source}")
                    except Exception as e:
                        log_lines.append(f"Step 1b: '{cmds.get('version')}' failed: {e}")

                    brief_raw = NetmikoService.clean_cli_output(
                        net_connect.send_command(cmds["lldp_brief"], read_timeout=settings.DEFAULT_TIMEOUT)
                    )
                    raw_parts.append(f"<{sysname}> {cmds['lldp_brief']}\n{brief_raw}")
                    brief_rows = cls.parse_lldp_brief(keep_lines(brief_raw, "lldp_brief", "Step 2"), parser)
                    log_lines.append(f"Step 2: '{cmds['lldp_brief']}' returned {len(brief_rows)} neighbor row(s)")

                    # Wrong vendor for this profile: the CLI rejected the LLDP command,
                    # or it returned nothing and the version command was rejected too
                    if cli_rejected(brief_raw) or (not brief_rows and cli_rejected(ver_raw)):
                        rejected_profiles += 1
                        if not is_last_profile:
                            log_lines.append(
                                f"Step 2b: [{cprof['name']}] commands rejected by this device, trying the next command profile"
                            )
                            continue

                    # Step 3: per-interface detail for every port that has a neighbor
                    intf_order: List[str] = []
                    details_by_intf: Dict[str, List[Dict[str, str]]] = {}
                    # A detail command without '{intf}' prints every port: run it once and
                    # split the neighbors by local port instead of once per port
                    shared_detail: Optional[Dict[str, List[Dict[str, str]]]] = None
                    if brief_rows and "{intf}" not in cmds["lldp_detail"]:
                        shared_detail = {}
                        try:
                            detail_raw = NetmikoService.clean_cli_output(
                                net_connect.send_command(cmds["lldp_detail"], read_timeout=settings.DEFAULT_TIMEOUT * 4)
                            )
                            raw_parts.append(f"<{sysname}> {cmds['lldp_detail']}\n{detail_raw}")
                            for r in cls.parse_detail(keep_lines(detail_raw, "lldp_detail", "Step 3"), parser):
                                shared_detail.setdefault(cls.intf_key(r["local_port"]), []).append(r)
                        except Exception as e:
                            log_lines.append(f"Step 3: '{cmds['lldp_detail']}' failed: {e}")
                    for row in brief_rows:
                        intf = row["local_port"]
                        if intf in details_by_intf:
                            continue
                        intf_order.append(intf)
                        if shared_detail is not None:
                            details_by_intf[intf] = shared_detail.get(cls.intf_key(intf), [])
                            log_lines.append(f"Step 3: '{cmds['lldp_detail']}' [{intf}] -> {len(details_by_intf[intf])} neighbor(s)")
                            continue
                        cmd = cmds["lldp_detail"].format(intf=intf) if "{intf}" in cmds["lldp_detail"] else cmds["lldp_detail"]
                        details: List[Dict[str, str]] = []
                        try:
                            detail_raw = NetmikoService.clean_cli_output(
                                net_connect.send_command(cmd, read_timeout=settings.DEFAULT_TIMEOUT)
                            )
                            raw_parts.append(f"<{sysname}> {cmd}\n{detail_raw}")
                            details = cls.parse_detail(
                                keep_lines(detail_raw, "lldp_detail", f"Step 3 [{intf}]"), parser, default_local_port=intf
                            )
                        except Exception as e:
                            log_lines.append(f"Step 3: '{cmd}' failed: {e}")
                        details_by_intf[intf] = details
                        log_lines.append(f"Step 3: '{cmd}' -> {len(details)} neighbor(s)")

                    # Step 4: ports without detail -> run the full LLDP detail command once
                    missing = [i for i in intf_order if not details_by_intf[i]]
                    same_as_detail = shared_detail is not None and cmds["lldp_full"] == cmds["lldp_detail"]
                    if (missing or not brief_rows) and not same_as_detail:
                        try:
                            full_raw = NetmikoService.clean_cli_output(
                                net_connect.send_command(cmds["lldp_full"], read_timeout=settings.DEFAULT_TIMEOUT * 4)
                            )
                            raw_parts.append(f"<{sysname}> {cmds['lldp_full']}\n{full_raw}")
                            full_rows = cls.parse_detail(keep_lines(full_raw, "lldp_full", "Step 4"), parser)
                            log_lines.append(
                                f"Step 4: {len(missing)} port(s) without detail, '{cmds['lldp_full']}' returned {len(full_rows)} neighbor(s)"
                            )
                            by_key: Dict[str, List[Dict[str, str]]] = {}
                            for r in full_rows:
                                by_key.setdefault(cls.intf_key(r["local_port"]), []).append(r)
                            if not brief_rows:
                                # Brief could not be parsed at all: use the full output as-is
                                for r in full_rows:
                                    if r["local_port"] not in details_by_intf:
                                        intf_order.append(r["local_port"])
                                        details_by_intf[r["local_port"]] = []
                                    details_by_intf[r["local_port"]].append(r)
                            for intf in missing:
                                details_by_intf[intf] = by_key.get(cls.intf_key(intf), [])
                        except Exception as e:
                            log_lines.append(f"Step 4: '{cmds['lldp_full']}' failed: {e}")

                    for intf in intf_order:
                        details = details_by_intf[intf]
                        if not details:
                            # Last resort: brief values for this port
                            log_lines.append(f"Step 4: {intf} has no detail anywhere, used brief values")
                            details = [
                                {
                                    "local_port": r["local_port"],
                                    "remote_device": r["remote_device"] or "N/A",
                                    "remote_port": r["remote_port"] or "N/A",
                                    "remote_ip": r.get("remote_ip", ""),
                                    "remote_model": "",
                                }
                                for r in brief_rows if r["local_port"] == intf
                            ]
                        for d in details:
                            neighbors.append({
                                "Local Device": sysname,
                                "Local Model": model,
                                "Local IP": local_ip,
                                "Local Port": d["local_port"],
                                "Remote Device": d["remote_device"],
                                "Remote Model": d.get("remote_model", ""),
                                "Remote Port": d["remote_port"],
                                "Remote IP": d.get("remote_ip", ""),
                                # Source of Remote Model, kept for 'Re-parse' after the model rules change
                                "Remote Description": d.get("remote_description", ""),
                            })

                    # Reaching here means the device accepted this profile's commands
                    # (a rejection already moved on above), so an empty neighbor list
                    # is a real answer - do not sweep the next vendor's commands.
                    log_lines.append(f"Step 5: Parsed {len(neighbors)} neighbors with [{cprof['name']}].")
                    break
        except Exception as e:
            # Netmiko messages start with blank lines and run over several; keep the log to one line
            status = f"Failed: {' '.join(str(e).split()) or type(e).__name__}"
            error = str(e)
            log_lines.append(f"ERROR: {status}")

        return {
            "hostname": sysname,
            "ip": local_ip,
            "model": model,
            # Source of the model, kept for 'Re-parse' after the model rules change
            "version_output": version_output,
            "command_profile": used_profile,
            "depth": depth,
            "status": status,
            "success": error is None,
            "error": error,
            "neighbors_found": len(neighbors),
            "neighbors": neighbors,
            # Every command profile was rejected: this driver / vendor is the wrong one
            "commands_rejected": bool(cmd_profiles) and rejected_profiles >= len(cmd_profiles),
            "log": "\n".join(log_lines),
            "raw_output": "\n\n".join(raw_parts),
            "execution_time_seconds": round(time.time() - start, 2),
        }

    @staticmethod
    def driver_plan(
        profiles: List[Dict[str, Any]], driver: Optional[str] = None, telnet: bool = False
    ) -> List[Any]:
        """
        [(driver, command profiles)] in the order to log in with.

        Known driver: that driver only, with the command profiles of its vendor ('cisco'
        matches cisco_ios / cisco_nxos, 'raisecom' raisecom_roap / raisecom_telnet, ...).
        A driver no profile is written for gets every profile, cycled on its one session.
        Unknown (driver None): one driver per profile parser, in command profile order.
        """
        from app.services.autodetect_service import AutoDetectService

        if driver:
            drv = driver.strip().lower()
            if telnet:
                drv = AutoDetectService.telnet_driver(drv)
            own = [p for p in profiles if p["parser"] in drv]
            return [(drv, own or profiles)]

        order: List[str] = []
        by_driver: Dict[str, List[Dict[str, Any]]] = {}
        for prof in profiles:
            drv = PARSER_DRIVERS.get(prof["parser"], prof["parser"])
            if telnet:
                drv = AutoDetectService.telnet_driver(drv)
            if drv not in by_driver:
                by_driver[drv] = []
                order.append(drv)
            by_driver[drv].append(prof)
        return [(drv, by_driver[drv]) for drv in order]

    @staticmethod
    def description_driver(description: str) -> Optional[str]:
        """Driver of the vendor an LLDP System Description names, None when it names none"""
        vendor = LldpService.detect_vendor(description or "", parser="-")
        return PARSER_DRIVERS.get(vendor) if vendor != "-" else None

    @classmethod
    def resolve_lldp_driver(cls, device: DeviceCredentials) -> Any:
        """
        (driver or None, where it came from, failure result or None).

        - a driver on the device (fleet row, or read from the LLDP description of a neighbor)
          is used as it is
        - 'unknown' -> None: the command profiles are swept
        - Auto Detect -> detected now (or the result of 'Detect Types' from the last few
          minutes); a vendor that cannot be named -> None (swept like 'unknown'). A dead
          host or rejected password returns a failure instead: logging in again with every
          driver cannot fix those and piles up failed logins.
        """
        from app.services.autodetect_service import AutoDetectService, AUTO_TYPES, is_driver
        dtype = (device.device_type or "").strip().lower()
        if is_driver(dtype):
            return dtype, "device type", None
        if dtype not in AUTO_TYPES:
            return None, "unknown", None
        try:
            detected, reason = AutoDetectService.detect_device_type(device)
        except Exception as e:
            return None, f"auto-detect error ({e})", None
        if is_driver(detected):
            return detected, f"auto-detect: {reason}", None
        if detected in ("unreachable", "auth_failed"):
            return None, detected, {
                "hostname": device.name or device.host, "ip": device.host, "model": "",
                "version_output": "", "command_profile": "", "status": f"Failed: {reason}",
                "success": False, "error": reason, "neighbors_found": 0, "neighbors": [],
                "commands_rejected": False, "raw_output": "", "execution_time_seconds": 0,
                "log": f"Auto-detect: {detected} - {reason}. No LLDP login attempted.",
            }
        return None, f"auto-detect could not name the vendor ({reason})", None

    @classmethod
    def collect_auto(
        cls, device: DeviceCredentials, depth: int = 0, command_profile_ids: Optional[List[str]] = None
    ) -> Dict[str, Any]:
        """Driver first (resolve_lldp_driver), then collect: one driver, or the sweep for unknown"""
        driver, source, failed = cls.resolve_lldp_driver(device)
        if depth > 0 and source == "device type":
            source = "the neighbor's LLDP system description"
        if failed:
            failed["depth"] = depth
            return failed
        result = cls.collect_device_sweep(device, depth, command_profile_ids, driver)
        note = (f"Driver [{driver}] from {source}" if driver
                else f"Driver unknown ({source}): sweeping the command profiles in order")
        result["log"] = f"{note}\n{result['log']}"
        return result

    @staticmethod
    def _other_driver_may_help(result: Dict[str, Any]) -> bool:
        """False for a wrong password or a dead host: logging in again with another driver
        cannot fix those, and on TACACS/RADIUS every extra failed login risks a lockout"""
        if result["success"]:
            return bool(result.get("commands_rejected"))
        # Logged in, but the prompt is not this driver's (Huawei driver on a 'Raisecom#' or
        # 'Switch#' prompt). Netmiko's message says "increase the read_timeout", which must
        # not be taken for a dead host: the next driver is the fix.
        if "pattern not detected" in (result.get("error") or "").lower():
            return True
        from app.services.job_service import JobService
        diag = JobService.format_failure_diagnostic(result.get("error") or "", host_ip=result.get("ip") or "")
        return not diag.startswith(("Authentication Failed", "Connection Timeout", "Network Unreachable", "Connection Refused"))

    @classmethod
    def collect_device_sweep(
        cls,
        device: DeviceCredentials,
        depth: int = 0,
        command_profile_ids: Optional[List[str]] = None,
        driver: Optional[str] = None,
    ) -> Dict[str, Any]:
        """
        Known `driver`: one login with it, running its own command profiles. Unknown
        (None): log in once per driver in command profile order and stop at the first one
        whose commands the device accepts, or at a wrong password / dead host.
        """
        from app.services.autodetect_service import AutoDetectService
        from app.services.command_profile_service import CommandProfileService
        profiles = CommandProfileService.resolve_ordered(command_profile_ids)
        plan = cls.driver_plan(profiles, driver, AutoDetectService.is_telnet(device))

        logs: List[str] = []
        result: Optional[Dict[str, Any]] = None
        for attempt, (drv, drv_profiles) in enumerate(plan, 1):
            result = cls.collect_device(device, depth, driver=drv, cmd_profiles=drv_profiles)
            logs.append(result["log"])
            if not cls._other_driver_may_help(result):
                break
            if attempt < len(plan):
                reason = "rejected the commands" if result["success"] else "could not be used"
                logs.append(f"--- Driver [{drv}] {reason}, retrying as [{plan[attempt][0]}] ---")
        if driver and result is not None and cls._other_driver_may_help(result):
            # The device type is followed as it is set: say so instead of a quiet
            # "Success, 0 neighbors" when it is the wrong one
            why = "rejected the LLDP commands" if result["success"] else "could not be used"
            msg = (f"Driver [{driver}] {why}: the device type may be wrong. Set it to Unknown to try every "
                   f"command profile, or pick the right driver")
            logs.append(f"--- {msg} ---")
            if result["success"]:
                result.update(success=False, status=f"Failed: {msg}", error=msg)
        if result is not None and len(logs) > 1:
            result["log"] = "\n".join(logs)
        return result if result is not None else cls.collect_device(device, depth, command_profile_ids)

    @classmethod
    def discover(
        cls,
        devices: List[DeviceCredentials],
        num_workers: Optional[int] = None,
        enable_tcp_scan: bool = False,
        scan_workers: int = 50,
        tcp_timeout: float = 1.5,
        recursive: bool = False,
        max_depth: int = 3,
        command_profile_ids: Optional[List[str]] = None,
    ) -> Dict[str, Any]:
        """
        Run collect_device over the seed devices. When recursive is on, the
        management IPs learned from LLDP are queued as the next wave (using the
        seed's credentials) until no new devices appear or max_depth is reached.
        """
        start = time.time()
        workers = max(1, min(100, num_workers or settings.DEFAULT_NUM_WORKERS))
        visited_ips: Set[str] = set()
        visited_names: Set[str] = set()
        hosts: List[Dict[str, Any]] = []
        seen_devices: Dict[str, List[Dict[str, Any]]] = {}
        # Log lines about the queue itself, appended to the first host of each wave
        queue_log: List[str] = []

        wave = []
        for d in devices:
            if d.host and d.host not in visited_ips:
                visited_ips.add(d.host)
                wave.append(d)
            # A fleet device is never re-visited through LLDP, whatever name it answers with
            if d.name:
                visited_names.add(cls.name_key(d.name))

        depth = 0
        while wave:
            results = [None] * len(wave)
            alive_indices = []

            if enable_tcp_scan:
                from app.services.lldp_scan_service import tcp_alive
                tcp_results = {}
                with ThreadPoolExecutor(max_workers=min(scan_workers, len(wave))) as tcp_exec:
                    tcp_futs = {
                        tcp_exec.submit(tcp_alive, dev.host, dev.port or settings.DEFAULT_SSH_PORT, tcp_timeout): i
                        for i, dev in enumerate(wave)
                    }
                    for fut in as_completed(tcp_futs):
                        tcp_results[tcp_futs[fut]] = fut.result()

                for i, dev in enumerate(wave):
                    if tcp_results.get(i):
                        alive_indices.append(i)
                    else:
                        port = dev.port or settings.DEFAULT_SSH_PORT
                        results[i] = {
                            "hostname": dev.name or dev.host,
                            "ip": dev.host,
                            "model": "",
                            "depth": depth,
                            "status": "UNREACHABLE",
                            "success": False,
                            "error": f"TCP/{port} port closed or no response within {tcp_timeout}s",
                            "neighbors_found": 0,
                            "neighbors": [],
                            "log": f"TCP/{port} closed or no response within {tcp_timeout}s (host down, ACL or firewall)",
                            "raw_output": "",
                            "execution_time_seconds": round(tcp_timeout, 2),
                        }
            else:
                alive_indices = list(range(len(wave)))

            if alive_indices:
                # Each device logs in with its driver: the fleet row's, the one auto-detect
                # names, or (neighbors) the one its LLDP description names. Unknown sweeps.
                with ThreadPoolExecutor(max_workers=min(workers, len(alive_indices))) as executor:
                    futures = {
                        executor.submit(cls.collect_auto, wave[i], depth, command_profile_ids): i
                        for i in alive_indices
                    }
                    for fut in as_completed(futures):
                        results[futures[fut]] = fut.result()

            next_wave = []
            for dev, res in zip(wave, results):
                hosts.append(res)
                if res["success"]:
                    visited_names.add(cls.name_key(res["hostname"]))
                    first = cls.link_same_device(res, seen_devices)
                    if first:
                        # One device on two IPs: keep this login's log, not its rows again
                        res["same_device_as"] = first["ip"]
                        res["status"] = "Same device"
                        res["detail"] = (f"Same device as {first['ip']} (same name, model and LLDP neighbors); "
                                         f"{res['neighbors_found']} neighbor row(s) not added again")
                        res["neighbors"] = []
                        res["neighbors_found"] = 0

            if recursive and depth < max_depth:
                # The whole wave is finished first, so every management IP it learned is
                # known before anything is queued: fleet devices and IPs already SSH'd
                # are dropped here, and the rest are SSH'd as the next wave.
                seen_here: Set[str] = set()
                total = skipped_fleet = skipped_dup = no_ip = 0
                for dev, res in zip(wave, results):
                    for n in res["neighbors"]:
                        total += 1
                        ip = (n.get("Remote IP") or "").strip()
                        name_key = cls.name_key(n.get("Remote Device") or "")
                        if not ip:
                            no_ip += 1
                            continue
                        # seen_here first, so a neighbor both sides report is counted
                        # as a duplicate rather than as an already-visited device
                        if ip in seen_here:
                            skipped_dup += 1
                            continue
                        if ip in visited_ips or name_key in visited_names:
                            skipped_fleet += 1
                            continue
                        seen_here.add(ip)
                        visited_ips.add(ip)
                        if name_key:
                            visited_names.add(name_key)
                        # The seed's credentials, but not its driver: the neighbor's own LLDP
                        # description names its vendor, else it is unknown and swept
                        n_driver = cls.description_driver(n.get("Remote Description") or n.get("Remote Model") or "")
                        next_wave.append(dev.copy(update={
                            "id": None,
                            "host": ip,
                            "name": n["Remote Device"],
                            "active_credential_name": None,
                            "device_type": n_driver or "unknown",
                        }))
                queue_log.append(
                    f"Depth {depth} -> {depth + 1}: {total} neighbor row(s) seen, {len(next_wave)} new IP(s) queued "
                    f"({skipped_fleet} already in the fleet or visited, {skipped_dup} duplicate, {no_ip} without a management IP)"
                )

            wave = next_wave
            depth += 1
            # Keep the queue decision visible in the logs of the wave that produced it
            if queue_log and hosts:
                hosts[-1]["log"] = f"{hosts[-1]['log']}\n{queue_log[-1]}"
                queue_log.clear()

        all_neighbors = [n for h in hosts for n in h["neighbors"]]
        cls.fill_remote_models(all_neighbors, hosts)
        # Devices, not IPs: a device that answered on two IPs counts once
        same_ips = sum(1 for h in hosts if h.get("same_device_as"))
        success = sum(1 for h in hosts if h["success"]) - same_ips
        return {
            "total_hosts": len(hosts),
            "success_hosts": success,
            "same_device_ips": same_ips,
            "failed_hosts": len(hosts) - success - same_ips,
            "total_lldp_rows": len(all_neighbors),
            "overall_time_seconds": round(time.time() - start, 2),
            "neighbors": all_neighbors,
            "hosts": hosts,
            "topology": cls.build_topology(all_neighbors, hosts),
        }

    # ---- One device answering on several IPs (Vlanif / MEth / loopback in several subnets) ----
    _SERIAL_LINE_RE = re.compile(
        r"^.*\b(?:serial\s*(?:number|no\.?)|system serial|esn|bar\s*code|(?:system|base|bridge)\s*mac(?:\s*address)?)\b.*$",
        re.I | re.M,
    )

    @classmethod
    def device_fingerprint(cls, host: Dict[str, Any]) -> Dict[str, Any]:
        """What tells two logins apart: model, the LLDP neighbors seen, serial / MAC lines"""
        return {
            "model": (host.get("model") or "").strip().lower(),
            "neighbors": frozenset(
                (cls.intf_key(n.get("Local Port") or ""), cls.name_key(n.get("Remote Device") or ""))
                for n in host.get("neighbors") or []
            ),
            "ids": frozenset(
                re.sub(r"\s+", " ", m.group(0)).strip().lower()
                for m in cls._SERIAL_LINE_RE.finditer(host.get("version_output") or "")
            ),
        }

    @classmethod
    def same_device(cls, a: Dict[str, Any], b: Dict[str, Any]) -> bool:
        """Two logged-in hosts are one device: same name, model and LLDP neighbors, and the
        same serial / MAC lines when both version outputs print them. A factory name shared
        by two boxes ('Raisecom', 'HUAWEI') is not enough on its own."""
        if cls.name_key(a.get("hostname") or "") != cls.name_key(b.get("hostname") or ""):
            return False
        fa, fb = cls.device_fingerprint(a), cls.device_fingerprint(b)
        if fa["model"] != fb["model"] or fa["neighbors"] != fb["neighbors"]:
            return False
        return not (fa["ids"] and fb["ids"]) or fa["ids"] == fb["ids"]

    @classmethod
    def link_same_device(cls, host: Dict[str, Any], seen: Dict[str, List[Dict[str, Any]]]) -> Optional[Dict[str, Any]]:
        """
        Register a logged-in host. Returns the host it is the same device as (that host
        gets this IP in 'other_ips'), else None. A different device with the same name is
        kept, and both carry 'same_name_as' so the name clash can be shown.
        """
        key = cls.name_key(host.get("hostname") or "")
        if not key:
            return None
        known = seen.setdefault(key, [])
        for first in known:
            if cls.same_device(first, host):
                first.setdefault("other_ips", []).append(host.get("ip") or "")
                return first
        for other in known:
            other.setdefault("same_name_as", []).append(host.get("ip") or "")
            host.setdefault("same_name_as", []).append(other.get("ip") or "")
        known.append(host)
        return None

    @staticmethod
    def version_from_raw(raw_output: str) -> str:
        """
        The 'display version' / 'show version' output inside a host's raw output
        (results collected before version_output was stored). Sections look like
        '<SW1> display version\n...'; the last one is from the profile that worked.
        """
        found = ""
        section = re.compile(r"^<[^<>\r\n]*> (.+)$", re.MULTILINE)
        heads = list(section.finditer(raw_output or ""))
        for i, m in enumerate(heads):
            if re.search(r"\bversion\b", m.group(1), re.I):
                end = heads[i + 1].start() if i + 1 < len(heads) else len(raw_output)
                body = raw_output[m.end():end]
                # Drop the '===== Command profile: X =====' line that starts the next profile
                found = re.split(r"^={5} Command profile: .*$", body, flags=re.MULTILINE)[0].strip()
        return found

    @classmethod
    def reparse(cls, neighbors: List[Dict[str, Any]], hosts: List[Dict[str, Any]]) -> Dict[str, Any]:
        """
        Re-run model extraction and role guessing on an existing result with the
        current model rules, without SSH. Hosts use their version output, neighbor
        rows their LLDP System description; rows without a source keep their model.
        """
        changed_hosts = changed_rows = 0
        for h in hosts:
            text = h.get("version_output") or cls.version_from_raw(h.get("raw_output") or "")
            if not text:
                continue
            model = cls.extract_model(text)
            if model and model != h.get("model"):
                h["model"] = model
                changed_hosts += 1
        host_models = {h.get("hostname"): h.get("model", "") for h in hosts if h.get("hostname")}
        for n in neighbors:
            before = (n.get("Local Model", ""), n.get("Remote Model", ""))
            if host_models.get(n.get("Local Device")):
                n["Local Model"] = host_models[n.get("Local Device")]
            desc = n.get("Remote Description") or ""
            if desc:
                n["Remote Model"] = cls.extract_model(desc) or n.get("Remote Model", "")
            if (n.get("Local Model", ""), n.get("Remote Model", "")) != before:
                changed_rows += 1
        cls.fill_remote_models(neighbors, hosts)
        return {
            "neighbors": neighbors,
            "hosts": hosts,
            "topology": cls.build_topology(neighbors, hosts),
            "changed_hosts": changed_hosts,
            "changed_rows": changed_rows,
        }

    @classmethod
    def build_topology(cls, neighbors: List[Dict[str, Any]], hosts: List[Dict[str, Any]]) -> Dict[str, Any]:
        """
        Nodes keyed by sysname (hostname, ip, model, role) and de-duplicated links:
        A:GE0/0/1 <-> B:GE0/0/2 reported from both sides becomes one link with confirmed=True.
        """
        nodes: Dict[str, Dict[str, Any]] = {}

        def ensure(name: str, ip: str = "", model: str = "") -> str:
            node = nodes.setdefault(name, {"id": name, "hostname": name, "ip": "", "model": "", "discovered": False})
            if ip and not node["ip"]:
                node["ip"] = ip
            if model and not node["model"]:
                node["model"] = model
            return name

        # SSH'd hosts first so their login IP and 'display version' model take priority
        for h in hosts:
            # A login that was the same device as an earlier IP adds nothing but that IP
            if h.get("success") and not h.get("same_device_as") and h.get("hostname"):
                node = nodes[ensure(h["hostname"], h.get("ip", ""), h.get("model", ""))]
                node["discovered"] = True
                if h.get("other_ips"):
                    node["other_ips"] = list(h["other_ips"])

        # Backup for FQDN names the rule above keeps: 'X.<domain>' is the same box as a known 'X'
        def canonical(name: str) -> str:
            if name in nodes or "." not in name:
                return name
            host = name.split(".", 1)[0]
            return host if host in nodes else name

        links: Dict[Any, Dict[str, Any]] = {}
        for n in neighbors:
            local = (n.get("Local Device") or "").strip()
            if not local:
                continue
            local = canonical(local)
            ensure(local, n.get("Local IP", ""), n.get("Local Model", ""))
            remote = (n.get("Remote Device") or "").strip()
            if not remote or remote.upper() == "N/A":
                remote = n.get("Remote IP") or f"Unknown ({local} {n.get('Local Port', '')})"
            else:
                remote = canonical(remote)
            ensure(remote, n.get("Remote IP", ""), n.get("Remote Model", ""))
            if remote == local:
                continue
            local_side = (local, cls.intf_key(n.get("Local Port", "")))
            remote_side = (remote, cls.intf_key(n.get("Remote Port", "")))
            key = tuple(sorted([local_side, remote_side]))
            if key in links:
                if links[key]["source"] != local:
                    links[key]["confirmed"] = True
                continue
            links[key] = {
                "source": local,
                "target": remote,
                "source_port": n.get("Local Port", ""),
                "target_port": n.get("Remote Port", ""),
                "confirmed": False,
            }

        degree: Dict[str, int] = {}
        for link in links.values():
            degree[link["source"]] = degree.get(link["source"], 0) + 1
            degree[link["target"]] = degree.get(link["target"], 0) + 1
        for node in nodes.values():
            info = cls.device_role_info(node["model"])
            node["role"] = info["role"]
            # Where the icon came from: a model rule, the built-in guess, or nothing
            node["role_source"] = info["source"]
            node["role_rule"] = info["rule_name"]
            node["degree"] = degree.get(node["id"], 0)

        # Model rules can also give firewall / server / wireless ...
        role_order = {"router": 0, "firewall": 1, "switch": 2, "wireless": 3, "server": 4, "unknown": 9}
        return {
            "nodes": sorted(nodes.values(), key=lambda x: (role_order.get(x["role"], 5), -x["degree"], x["hostname"])),
            "links": list(links.values()),
        }

    @staticmethod
    def build_excel(neighbors: List[Dict[str, Any]], hosts: List[Dict[str, Any]]) -> bytes:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()

        def write_sheet(ws, headers, rows):
            ws.append(headers)
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
            for r in rows:
                ws.append([
                    (str(r.get(h, "")) if r.get(h) is not None else "")[:EXCEL_CELL_LIMIT]
                    for h in headers
                ])
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for i, h in enumerate(headers, 1):
                width = max([len(h)] + [min(len(str(r.get(h, "") or "")), 60) for r in rows[:500]])
                ws.column_dimensions[get_column_letter(i)].width = width + 2

        # Sheet 1: all LLDP entries
        ws1 = wb.active
        ws1.title = "LLDP_Inventory"
        write_sheet(
            ws1,
            ["Local Device", "Local Model", "Local IP", "Local Port", "Remote Device", "Remote Model", "Remote Port", "Remote IP"],
            neighbors,
        )

        # Sheet 2: host list & status (shows which hosts have no neighbors)
        ws2 = wb.create_sheet("Execution_Summary")
        write_sheet(
            ws2,
            ["Hostname", "IP Address", "Other IPs", "Model", "Depth", "Status", "Neighbors Found", "Time (s)"],
            [
                {
                    "Hostname": h.get("hostname"),
                    "IP Address": h.get("ip"),
                    "Other IPs": ", ".join(h.get("other_ips") or []),
                    "Model": h.get("model", ""),
                    "Depth": h.get("depth", 0),
                    "Status": h.get("status"),
                    "Neighbors Found": h.get("neighbors_found", 0),
                    "Time (s)": h.get("execution_time_seconds"),
                }
                # One row per device: its other IPs are in 'Other IPs'
                for h in hosts if not h.get("same_device_as")
            ],
        )

        # Sheet 3: per-host log + raw terminal output (replaces logs/<host>.log)
        ws3 = wb.create_sheet("Raw_Logs")
        write_sheet(
            ws3,
            ["Hostname", "IP Address", "Log", "Raw Output"],
            [
                {
                    "Hostname": h.get("hostname"),
                    "IP Address": h.get("ip"),
                    "Log": h.get("log"),
                    "Raw Output": h.get("raw_output") or "No Data",
                }
                for h in hosts
            ],
        )

        stream = io.BytesIO()
        wb.save(stream)
        return stream.getvalue()
